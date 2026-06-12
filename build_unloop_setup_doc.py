from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path("/Users/yingyao/Downloads/VampNet/Unloop_Setup_and_Model_Guide.docx")


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_text(cell, text, bold=False, color=None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(text)
    run.bold = bold
    run.font.name = "Calibri"
    run.font.size = Pt(9.5)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_table_widths(table, widths):
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = Inches(width)


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        set_cell_text(hdr[i], h, bold=True, color="0B2545")
        set_cell_shading(hdr[i], "E8EEF5")
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], value)
    if widths:
        set_table_widths(table, widths)
    doc.add_paragraph()
    return table


def add_code(doc, text):
    p = doc.add_paragraph()
    p.style = "Code Block"
    for line in text.strip("\n").split("\n"):
        if line:
            p.add_run(line)
        p.add_run("\n")
    return p


def add_note(doc, title, body):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    cell = table.cell(0, 0)
    set_cell_shading(cell, "F4F6F9")
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(title)
    r.bold = True
    r.font.color.rgb = RGBColor(31, 58, 95)
    r.font.size = Pt(10)
    p2 = cell.add_paragraph(body)
    p2.paragraph_format.space_after = Pt(0)
    for run in p2.runs:
        run.font.size = Pt(9.5)
    doc.add_paragraph()


def add_kv_table(doc, rows):
    return add_table(doc, ["Item", "Value"], rows, widths=[1.75, 4.75])


def add_bullet(doc, text):
    p = doc.add_paragraph(text, style="List Bullet")
    p.paragraph_format.space_after = Pt(4)
    return p


def add_step(doc, text):
    p = doc.add_paragraph(text, style="List Number")
    p.paragraph_format.space_after = Pt(4)
    return p


def add_h(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    return p


def setup_styles(doc):
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    title = styles["Title"]
    title.font.name = "Calibri"
    title.font.size = Pt(22)
    title.font.bold = True
    title.font.color.rgb = RGBColor(11, 37, 69)
    title.paragraph_format.space_after = Pt(4)

    subtitle = styles.add_style("Guide Subtitle", 1)
    subtitle.font.name = "Calibri"
    subtitle.font.size = Pt(11)
    subtitle.font.color.rgb = RGBColor(85, 85, 85)
    subtitle.paragraph_format.space_after = Pt(12)

    for name, size, color, before, after in [
        ("Heading 1", 16, "2E74B5", 18, 10),
        ("Heading 2", 13, "2E74B5", 14, 7),
        ("Heading 3", 12, "1F4D78", 10, 5),
    ]:
        s = styles[name]
        s.font.name = "Calibri"
        s.font.size = Pt(size)
        s.font.color.rgb = RGBColor.from_string(color)
        s.paragraph_format.space_before = Pt(before)
        s.paragraph_format.space_after = Pt(after)

    code = styles.add_style("Code Block", 1)
    code.font.name = "Courier New"
    code.font.size = Pt(8.5)
    code.paragraph_format.left_indent = Inches(0.18)
    code.paragraph_format.right_indent = Inches(0.05)
    code.paragraph_format.space_before = Pt(2)
    code.paragraph_format.space_after = Pt(8)
    code.paragraph_format.line_spacing = 1.05

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = footer.add_run("Unloop setup guide")
    r.font.size = Pt(8)
    r.font.color.rgb = RGBColor(85, 85, 85)


def build():
    doc = Document()
    setup_styles(doc)

    title = doc.add_paragraph(style="Title")
    title.add_run("Unloop Setup and Model Weights Guide")
    subtitle = doc.add_paragraph(
        "For /Users/yingyao/Downloads/VampNet on macOS with Max 9 and vampnet_env",
        style="Guide Subtitle",
    )
    subtitle.alignment = WD_ALIGN_PARAGRAPH.LEFT

    add_note(
        doc,
        "Current working state",
        "The local VampNet server runs at http://127.0.0.1:7860. The main performance patch is "
        "/Users/yingyao/Downloads/VampNet/unloop-main/unloop.maxpat. The Max-side default model is currently set to percussion.",
    )

    add_h(doc, "1. What This Setup Does")
    doc.add_paragraph(
        "Unloop is a phrase-capture workflow. Max records a short loop, writes it to disk, launches a Python command through the shell external, "
        "sends the audio to the local Gradio/VampNet server, waits for the generated file, then loads the output back into Max."
    )
    add_bullet(doc, "It is not a zero-latency footstep-by-footstep effect.")
    add_bullet(doc, "For live tap dance, treat it as call-and-response: tap phrase, click unloop, wait, then perform with the generated loop.")
    add_bullet(doc, "Use low sampling steps for performance tests; use higher steps for final renders.")

    add_h(doc, "2. Folder Map and Assumptions")
    add_kv_table(
        doc,
        [
            ("Workspace root", "/Users/yingyao/Downloads/VampNet"),
            ("VampNet app", "/Users/yingyao/Downloads/VampNet/vampnet-main"),
            ("Main Max patch", "/Users/yingyao/Downloads/VampNet/unloop-main/unloop.maxpat"),
            ("Max subpatch", "/Users/yingyao/Downloads/VampNet/unloop-main/vamper.maxpat"),
            ("Python environment", "/Users/yingyao/anaconda3/envs/vampnet_env"),
            ("Python executable", "/Users/yingyao/anaconda3/envs/vampnet_env/bin/python"),
            ("Local API", "http://127.0.0.1:7860/"),
            ("Generated audio folder", "/Users/yingyao/Downloads/VampNet/unloop-main/audio"),
        ],
    )

    add_h(doc, "3. Install and Verify Max Externals")
    doc.add_paragraph("The patch depends on the Max externals shell and karma~. Install both before opening the patch.")
    add_step(doc, "Check the active Max package folders.")
    add_code(
        doc,
        r"""
find "$HOME/Documents/Max 9/Packages" "$HOME/Library/Application Support/Cycling '74/Max 9/Packages" \
  -iname 'karma~.mxo' -o -iname 'shell.mxo'
""",
    )
    add_step(doc, "If shell is missing, install the shell package into Max's package folder.")
    add_code(
        doc,
        r"""
mkdir -p "$HOME/Library/Application Support/Cycling '74/Max 9/Packages"
# Download shell.zip from:
# https://github.com/jeremybernstein/shell/releases/download/1.0b3/shell.zip
# Unzip so this exists:
# ~/Library/Application Support/Cycling '74/Max 9/Packages/shell/externals/shell.mxo
""",
    )
    add_step(doc, "If karma~ is missing, install karma 1.6.1 into Max's package folder.")
    add_code(
        doc,
        r"""
# Expected final path:
# ~/Library/Application Support/Cycling '74/Max 9/Packages/karma/externals/karma~.mxo
# Package info should report version 1.6.1.
""",
    )
    add_step(doc, "If Max still reports 'shell: No such object', copy the package into Documents/Max 9/Packages as well, then restart Max.")
    add_code(
        doc,
        r"""
ditto "$HOME/Library/Application Support/Cycling '74/Max 9/Packages/shell" \
      "$HOME/Documents/Max 9/Packages/shell"
ditto "$HOME/Library/Application Support/Cycling '74/Max 9/Packages/karma" \
      "$HOME/Documents/Max 9/Packages/karma"
""",
    )
    add_note(
        doc,
        "Verification target",
        "When Max is running, lsof should show Max loading shell.mxo and karma~.mxo. If not, restart Max after installing packages.",
    )

    add_h(doc, "4. Prepare the Python Environment")
    doc.add_paragraph("Use the existing conda environment named vampnet_env. Install the dependencies that the app and Max bridge need.")
    add_code(
        doc,
        r"""
/Users/yingyao/anaconda3/envs/vampnet_env/bin/pip install python-osc
/Users/yingyao/anaconda3/envs/vampnet_env/bin/pip install -e /Users/yingyao/Downloads/VampNet/wavebeat-main
/Users/yingyao/anaconda3/envs/vampnet_env/bin/pip install -e /Users/yingyao/Downloads/VampNet/unloop-main
""",
    )
    add_step(doc, "Verify the environment's Python path.")
    add_code(doc, "/Users/yingyao/anaconda3/envs/vampnet_env/bin/python --version")

    add_h(doc, "5. Run the Local VampNet Server")
    doc.add_paragraph("Start the local Gradio server from the VampNet app folder. Keep this process running while using Max.")
    add_code(
        doc,
        r"""
cd /Users/yingyao/Downloads/VampNet/vampnet-main
PYTHONUNBUFFERED=1 /Users/yingyao/anaconda3/envs/vampnet_env/bin/python app.py
""",
    )
    doc.add_paragraph("Successful startup prints a line like:")
    add_code(doc, "* Running on local URL:  http://127.0.0.1:7860")
    add_step(doc, "Verify the server is listening.")
    add_code(doc, "lsof -nP -iTCP:7860 -sTCP:LISTEN")
    add_step(doc, "Verify Gradio endpoints are visible.")
    add_code(doc, "curl -s --noproxy '*' http://127.0.0.1:7860/gradio_api/info | jq -r '.named_endpoints | keys[]'")

    add_h(doc, "6. Required Code and Patch Changes")
    add_table(
        doc,
        ["File", "Required change"],
        [
            (
                "vampnet-main/app.py",
                "Run on 127.0.0.1:7860; keep localhost out of proxy routing; allow cached model names when Hugging Face is slow.",
            ),
            (
                "unloop-main/vamp.py",
                "Add a wrapper script that normalizes Max shell arguments before calling argbind.",
            ),
            (
                "unloop-main/unloop/unloop.py",
                "Use Gradio Client with trust_env=False, send input with handle_file, call api_name='/vamp_1', and default beat_mask_ms to 0.",
            ),
            (
                "unloop-main/vamper.maxpat",
                "Set Python path, server URL, base folder, absolute vamp.py command path, and ckpt-name/model menu.",
            ),
            (
                "unloop-main/unloop.maxpat",
                "Point dependencies to unloop-main and replace missing min.pan~ with local pan~.",
            ),
            (
                "unloop-main/pan~.maxpat",
                "Copy local pan abstraction into unloop-main so Max does not need min.pan~.",
            ),
        ],
        widths=[2.15, 4.35],
    )

    add_h(doc, "6.1 app.py")
    doc.add_paragraph("Use these server-side settings so browser, Max, and the Gradio client agree on the same local endpoint.")
    add_code(
        doc,
        r"""
os.environ["NO_PROXY"] = "localhost,127.0.0.1,::1," + os.environ.get("NO_PROXY", "")
os.environ.setdefault("HF_HUB_ETAG_TIMEOUT", "60")
os.environ.setdefault("HF_HUB_DOWNLOAD_TIMEOUT", "300")

# At launch:
demo.launch(server_name="127.0.0.1", server_port=7860)
""",
    )

    add_h(doc, "6.2 unloop-main/vamp.py")
    doc.add_paragraph("This file is the Max-friendly entry point. It splits quoted Max shell atoms like '--audio_path /path/file.wav' and drops legacy beatmask/downbeats atoms.")
    add_code(
        doc,
        r"""
import argbind
import shlex
import sys

from unloop.unloop import Vamp


def _normalize_max_shell_args():
    normalized = [sys.argv[0]]
    for arg in sys.argv[1:]:
        if arg in {"beatmask", "downbeats"}:
            continue
        if arg.startswith("--") and " " in arg:
            normalized.extend(shlex.split(arg))
        else:
            normalized.append(arg)
    sys.argv[:] = normalized


if __name__ == "__main__":
    _normalize_max_shell_args()
    args = argbind.parse_args()
    with argbind.scope(args):
        Vamp()
""",
    )

    add_h(doc, "6.3 unloop-main/unloop/unloop.py")
    doc.add_paragraph("The important behaviors are local API access, the Gradio 5 endpoint, and disabling beat-mask extraction by default.")
    add_code(
        doc,
        r"""
Client(servername, verbose=False, httpx_kwargs={"trust_env": False})

client.submit(
    input_audio=handle_file(audio_path),
    model_choice=checkpoint_name,
    beat_mask_ms=beat_mask_ms,
    api_name="/vamp_1",
)

# Safe default:
beat_mask_ms: int = 0
""",
    )
    add_note(
        doc,
        "Why beat_mask_ms is 0",
        "A nonzero beat mask can trigger WaveBeat. On this MPS setup, that path produced a CPU/MPS tensor mismatch. Leave it at 0 unless you explicitly fix and test WaveBeat.",
    )

    add_h(doc, "6.4 Max Patch Changes")
    add_table(
        doc,
        ["Control or dependency", "Value"],
        [
            ("Python path", "/Users/yingyao/anaconda3/envs/vampnet_env/bin/python"),
            ("API URL", "http://127.0.0.1:7860/"),
            ("Base folder", "/Users/yingyao/Downloads/VampNet/unloop-main"),
            ("Shell command script", "/Users/yingyao/Downloads/VampNet/unloop-main/vamp.py"),
            ("Main patch", "/Users/yingyao/Downloads/VampNet/unloop-main/unloop.maxpat"),
            ("Subpatch", "/Users/yingyao/Downloads/VampNet/unloop-main/vamper.maxpat"),
            ("Pan object", "Use local pan~.maxpat instead of missing min.pan~.mxo"),
            ("Default model", "percussion, with default as fallback"),
        ],
        widths=[2.05, 4.45],
    )
    doc.add_paragraph("After changing JSON or patcher settings, close and reopen unloop.maxpat so Max reloads pattr storage and dependency paths.")

    add_h(doc, "7. Operating Unloop in Max")
    add_step(doc, "Start app.py and confirm http://127.0.0.1:7860 is listening.")
    add_step(doc, "Open /Users/yingyao/Downloads/VampNet/unloop-main/unloop.maxpat.")
    add_step(doc, "In Max, open Audio Status and choose your input mic/interface and output device.")
    add_step(doc, "Turn Max audio on and raise input-gain until tap input is present without clipping.")
    add_step(doc, "Record a short phrase, ideally 3 to 8 seconds.")
    add_step(doc, "Click unloop once. Do not click repeatedly; each click can queue another job.")
    add_step(doc, "Wait for STATUS.DONE or for the output waveform/file to appear.")
    add_step(doc, "Bring up the wet gain and perform with the generated loop. Use feedback only after the first output has loaded.")
    add_note(
        doc,
        "Live timing",
        "numsteps=36 can take several minutes. For live rehearsals, set numsteps to 4, 8, or 14. Use 36 only when quality matters more than speed.",
    )

    add_h(doc, "8. Model Weights: Install, Verify, Switch")
    doc.add_paragraph("The default model uses the root files coarse.pth, c2f.pth, codec.pth, and wavebeat.pth. Fine-tuned models live under models/vampnet/loras/MODEL_NAME/.")
    add_kv_table(
        doc,
        [
            ("Fine-tuned coarse file", "models/vampnet/loras/MODEL_NAME/coarse.pth"),
            ("Fine-tuned c2f file", "models/vampnet/loras/MODEL_NAME/c2f.pth"),
            ("Currently installed LoRAs", "percussion and machines"),
            ("Best tap choice", "percussion"),
            ("Stable fallback", "default"),
        ],
    )

    add_h(doc, "8.1 Download a Fine-Tuned Model")
    doc.add_paragraph("Use the Hugging Face CLI. Replace MODEL_NAME with percussion, machines, orchestral, choir, and so on.")
    add_code(
        doc,
        r"""
cd /Users/yingyao/Downloads/VampNet/vampnet-main

HF_HUB_ETAG_TIMEOUT=60 HF_HUB_DOWNLOAD_TIMEOUT=600 \
/Users/yingyao/anaconda3/bin/hf download hugggof/vampnet \
  --include 'loras/MODEL_NAME/coarse.pth' \
  --include 'loras/MODEL_NAME/c2f.pth' \
  --local-dir /Users/yingyao/Downloads/VampNet/vampnet-main/models/vampnet \
  --max-workers 1
""",
    )
    doc.add_paragraph("Example: download machines.")
    add_code(
        doc,
        r"""
HF_HUB_ETAG_TIMEOUT=60 HF_HUB_DOWNLOAD_TIMEOUT=600 \
/Users/yingyao/anaconda3/bin/hf download hugggof/vampnet \
  --include 'loras/machines/coarse.pth' \
  --include 'loras/machines/c2f.pth' \
  --local-dir /Users/yingyao/Downloads/VampNet/vampnet-main/models/vampnet \
  --max-workers 1
""",
    )
    doc.add_paragraph("Verify the model files.")
    add_code(
        doc,
        r"""
ls -lh models/vampnet/loras/MODEL_NAME/coarse.pth \
       models/vampnet/loras/MODEL_NAME/c2f.pth
""",
    )
    add_note(
        doc,
        "Network behavior",
        "If a model is not downloaded before generation, Gradio may sit on STATUS.PROCESSING and then error while trying to fetch weights. Download model weights before performance.",
    )

    add_h(doc, "8.2 Switch Models in the Browser")
    add_step(doc, "Refresh http://127.0.0.1:7860 after downloading weights.")
    add_step(doc, "Choose the model from the model choice dropdown.")
    add_step(doc, "Use low sampling steps for the first test so the model loads and runs quickly.")
    add_step(doc, "If the browser shows error, check the server log for missing coarse.pth or c2f.pth.")

    add_h(doc, "8.3 Switch Models in unloop.maxpat")
    doc.add_paragraph("The Max presentation view does not expose the model menu directly. The reliable method is to update vamper.maxpat and then reopen unloop.maxpat.")
    add_step(doc, "Close unloop.maxpat in Max.")
    add_step(doc, "Open /Users/yingyao/Downloads/VampNet/unloop-main/vamper.maxpat as text or edit it in Max patching mode.")
    add_step(doc, "Find the pattr restore entry named ckpt-name and set it to the desired model.")
    add_code(doc, '"ckpt-name" : [ "percussion" ],')
    add_step(doc, "Make sure the hidden model menu items include that model name.")
    add_code(
        doc,
        r'''"items" : [ "default", ",", "percussion", ",", "machines", ",", "orchestral", ... ]'''
    )
    add_step(doc, "Save vamper.maxpat, then reopen unloop.maxpat.")
    add_step(doc, "Run a short test. The Max log or server log should show --checkpoint_name MODEL_NAME / model_choice: MODEL_NAME.")

    add_h(doc, "9. Troubleshooting")
    add_table(
        doc,
        ["Symptom", "Likely cause", "Fix"],
        [
            (
                "STATUS.PROCESSING for a long time",
                "Generation is still running, a job is queued, or the model is downloading.",
                "Check app.py log and ps. Do not click unloop again. Use lower numsteps.",
            ),
            (
                "Gradio shows error after processing",
                "Missing model weights or Hugging Face timeout.",
                "Download coarse.pth and c2f.pth for the selected model, then restart app.py.",
            ),
            (
                "shell: No such object",
                "Max did not find shell.mxo.",
                "Install shell into the active Max package folder and restart Max.",
            ),
            (
                "karma~ missing",
                "Max did not find karma~.mxo.",
                "Install karma 1.6.1 and restart Max.",
            ),
            (
                "No output file appears",
                "Python job still running, crashed, or output path mismatch.",
                "Check Max log, app.py terminal, and unloop-main/audio modification times.",
            ),
            (
                "Clicking unloop does nothing",
                "Patch not reloaded, no recorded audio, or shell command did not launch.",
                "Reopen unloop.maxpat, record a phrase first, and watch Max Console.",
            ),
            (
                "Beat mask crash",
                "WaveBeat path triggered on MPS/CPU mismatch.",
                "Keep beat_mask_ms at 0.",
            ),
        ],
        widths=[1.75, 2.2, 2.55],
    )

    add_h(doc, "10. Useful Checks During Performance")
    add_code(
        doc,
        r"""
# Is the server running?
lsof -nP -iTCP:7860 -sTCP:LISTEN

# Is Max currently running a Python generation command?
ps -axo pid,ppid,stat,etime,pcpu,command | rg 'unloop\.py|vamp\.py|app\.py|Max.app'

# Did a new file arrive?
ls -lt /Users/yingyao/Downloads/VampNet/unloop-main/audio | head -20

# What does Max think happened?
tail -n 120 "$HOME/Library/Application Support/Cycling '74/Max 9/Logs/Max.log"
""",
    )

    add_h(doc, "11. Recommended Tap Dance Settings")
    add_table(
        doc,
        ["Use case", "Model", "numsteps", "Notes"],
        [
            ("Fast live rehearsal", "percussion or default", "4 to 8", "Fastest way to test the loop workflow."),
            ("Demo recording", "percussion", "8 to 14", "Good balance of time and texture."),
            ("Final render", "percussion", "36", "Slower, but more complete generation."),
            ("Emergency fallback", "default", "8 to 14", "Use if fine-tuned model loading fails."),
        ],
        widths=[1.55, 1.45, 1.2, 2.3],
    )

    add_note(
        doc,
        "Performance rule",
        "Tap phrase -> click unloop once -> wait for DONE -> perform with the generated loop. Avoid sending another job while STATUS.PROCESSING is active.",
    )

    doc.save(OUT)
    return OUT


if __name__ == "__main__":
    print(build())
